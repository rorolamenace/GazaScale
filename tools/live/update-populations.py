#!/usr/bin/env python3
"""Refreshes commune populations from their official sources and publishes the changed files on the
live-data branch, so the site picks them up without a redeploy.

Usage: python3 tools/live/update-populations.py DIST_DIR LIVE_DIR
  DIST_DIR: the site (dist/) with the commune files of the release;
  LIVE_DIR: a checkout of the live-data branch (manifest.json is updated there).
Each source is tried for a newer year than the one already published; nothing is written otherwise.
Geometry never changes here: new boundaries need a release of the site.
Needs: openpyxl (Israel, Japan, Australia). Prints a summary; exits 0 even when a source is unreachable (it is retried next run).
"""
import csv, io, json, os, sys, glob, datetime, urllib.request, urllib.error

DIST, LIVE = sys.argv[1], sys.argv[2]
UA = {'User-Agent': 'GazaScale live data (https://gazascale.org)'}
NOW = datetime.date.today().year
manifest_path = os.path.join(LIVE, 'manifest.json')
manifest = json.load(open(manifest_path))
datasets = manifest.setdefault('datasets', {})
# Years built into the release: a source must be newer than this (or than what live-data already has).
BUILT_IN = {'eu': 2024, 'it': 2024, 'me': 2024, 'br': 2024, 'us': 2024, 'ca': 2025, 'jp': 2026, 'au': 2025}

def get(url, binary=False):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=300) as r:
        data = r.read()
    return data if binary else data.decode('utf-8-sig', 'replace')

def exists(url):
    try:
        urllib.request.urlopen(urllib.request.Request(url, headers=UA, method='HEAD'), timeout=60).close()
        return True
    except (urllib.error.URLError, TimeoutError):
        return False

def published(key):
    return (datasets.get('_years') or {}).get(key, BUILT_IN[key])

def apply(dir_name, values, year, key, source, only=None):
    """Writes every file of dist/<dir_name> (or only those matching `only`) with the new populations."""
    files = sorted(glob.glob(os.path.join(DIST, dir_name, '*.json')))
    if only: files = [f for f in files if only(os.path.basename(f))]
    written, changed = [], 0
    os.makedirs(os.path.join(LIVE, dir_name), exist_ok=True)
    for f in files:
        name = os.path.basename(f)
        if name in ('index.json', 'countries.json', 'places.json'): continue
        live_copy = os.path.join(LIVE, dir_name, name)
        data = json.load(open(live_copy if os.path.exists(live_copy) else f))
        hit = False
        for feat in data.get('features', []):
            p = feat['properties']; v = values.get(p.get('code'))
            if v and v > 0:
                if p.get('population') != v: changed += 1
                p['population'] = v; p['populationYear'] = year; p.pop('estimated', None); hit = True
        if hit:
            json.dump(data, open(live_copy, 'w'), separators=(',', ':'), ensure_ascii=False)
            written.append(name)
    entry = datasets.setdefault(dir_name, {})
    entry['year'] = max(entry.get('year', 0), year)
    if only or dir_name == 'wd':
        entry['files'] = sorted(set(entry.get('files', [])) | set(written))
    else:
        # The client reads every file of the folder from live-data: copy the untouched ones too.
        for f in files:
            name = os.path.basename(f); dst = os.path.join(LIVE, dir_name, name)
            if name not in written and not os.path.exists(dst) and name not in ('index.json', 'countries.json', 'places.json'):
                open(dst, 'wb').write(open(f, 'rb').read())
        entry.pop('files', None)
    entry.setdefault('sources', {})[key] = source
    datasets.setdefault('_years', {})[key] = year
    print(f'{key}: {year}, {changed} populations changed in {len(written)} files ({source})')
    return True

def eurostat():
    """Eurostat GISCO LAU: Europe (eu/) and Italy (it/)."""
    for year in range(NOW, max(published('eu'), published('it')), -1):
        url = f'https://gisco-services.ec.europa.eu/distribution/v2/lau/csv/LAU_RG_01M_{year}_4326.csv'
        if not exists(url): continue
        eu, it = {}, {}
        for r in csv.DictReader(io.StringIO(get(url))):
            try: v = round(float(r.get(f'POP_{year}') or 0))
            except ValueError: continue
            if v <= 0: continue
            gid = r['GISCO_ID']
            if r['CNTR_CODE'] == 'IT': it['IT-' + gid[3:]] = v
            else: eu['EU-' + gid] = v
        if len(eu) < 10000: continue  # population column not filled yet for this edition
        if year > published('eu'): apply('eu', eu, year, 'eu', url)
        if year > published('it') and len(it) > 7000: apply('it', it, year, 'it', url)
        return

def israel():
    """Israel CBS localities file (bycodeYYYY.xlsx): Israeli localities of me/IL.json and me/PS.json."""
    import openpyxl
    for year in range(NOW, published('me'), -1):
        url = f'https://www.cbs.gov.il/he/publications/doclib/2019/ishuvim/bycode{year}.xlsx'
        try: data = get(url, binary=True)
        except urllib.error.URLError: continue
        if len(data) < 100000: continue
        rows = list(openpyxl.load_workbook(io.BytesIO(data), read_only=True).worksheets[0].iter_rows(values_only=True))
        head = {k: i for i, k in enumerate(rows[0])}
        col = head.get(f'סך הכל אוכלוסייה {year}'); code = head.get('סמל יישוב')
        if col is None or code is None: continue
        values = {f'IL-{r[code]}': int(r[col]) for r in rows[1:] if r[code] and isinstance(r[col], (int, float)) and r[col] > 0}
        if len(values) > 1000: apply('me', values, year, 'me', url)
        return

def brazil():
    """IBGE population estimates by municipality (SIDRA table 6579), in wd/BR-*.json."""
    try: rows = json.loads(get('https://apisidra.ibge.gov.br/values/t/6579/n6/all/v/all/p/last'))
    except Exception as e: print('br: unreachable', e); return
    year = int(rows[1]['D3C'])
    if year <= published('br'): return
    values = {'WD-BR-' + r['D1C']: int(r['V']) for r in rows[1:] if r['V'].isdigit()}
    if len(values) > 5000:
        apply('wd', values, year, 'br', 'https://sidra.ibge.gov.br/tabela/6579', only=lambda n: n.startswith('BR-'))

def united_states():
    """US Census subcounty estimates: official for towns and townships, county growth elsewhere."""
    for year in range(NOW, published('us'), -1):
        url = f'https://www2.census.gov/programs-surveys/popest/datasets/2020-{year}/cities/totals/sub-est{year}.csv'
        if not exists(url): continue
        rows = list(csv.DictReader(io.StringIO(get(url).encode('latin-1', 'ignore').decode('latin-1'))))
        prev = published('us')
        county = {r['STATE'] + r['COUNTY']: (int(r[f'POPESTIMATE{year}']), int(r[f'POPESTIMATE{prev}']) or 1) for r in rows if r['SUMLEV'] == '050'}
        mcd = {r['STATE'] + r['COUNTY'] + r['COUSUB']: int(r[f'POPESTIMATE{year}']) for r in rows if r['SUMLEV'] == '061'}
        values = {}
        for f in glob.glob(os.path.join(DIST, 'us', '*.json')):
            live_copy = os.path.join(LIVE, 'us', os.path.basename(f))
            for feat in json.load(open(live_copy if os.path.exists(live_copy) else f))['features']:
                p = feat['properties']; g = p['code'][3:]
                if g in mcd: values[p['code']] = mcd[g]
                elif g[:5] in county and p.get('population'): values[p['code']] = round(p['population'] * county[g[:5]][0] / county[g[:5]][1])
        if len(values) > 30000: apply('us', values, year, 'us', url)
        return

def canada():
    """Statistics Canada table 17-10-0155 (July 1 estimates by census subdivision), in wd/CA-*.json."""
    import zipfile
    data = zipfile.ZipFile(io.BytesIO(get('https://www150.statcan.gc.ca/n1/tbl/csv/17100155-eng.zip', binary=True)))
    rows = list(csv.DictReader(io.StringIO(data.read('17100155.csv').decode('utf-8-sig'))))
    year = max(int(r['REF_DATE']) for r in rows)
    if year <= published('ca'): return
    values = {'WD-CA-' + r['DGUID'][9:]: int(r['VALUE']) for r in rows
              if r['REF_DATE'] == str(year) and r['DGUID'].startswith('2021A0005') and r['VALUE'].isdigit()}
    if len(values) > 4000:
        apply('wd', values, year, 'ca', 'https://www150.statcan.gc.ca/t1/tbl1/en/tv.action?pid=1710015501', only=lambda n: n.startswith('CA-'))

def japan():
    """Japanese resident register at January 1 (soumu.go.jp, municipalities, all residents), in wd/JP-*.json."""
    import openpyxl, re
    page_url = 'https://www.soumu.go.jp/main_sosiki/jichi_gyousei/daityo/jinkou_jinkoudoutai-setaisuu.html'
    page = get(page_url, binary=True).decode('shift_jis', 'replace')
    # The page lists one workbook per table: the right one is the all-residents (総計) population by municipality (市区町村別).
    links = [m.group(1) for m in re.finditer(r'href="([^"]+\.xlsx)">([^<]*)', page)
             if '総計' in m.group(2) and '市区町村別' in m.group(2) and '人口・世帯数' in m.group(2)]
    if not links: print('jp: workbook link not found'); return
    rows = list(openpyxl.load_workbook(io.BytesIO(get('https://www.soumu.go.jp' + links[0], binary=True)), read_only=True).worksheets[0].iter_rows(values_only=True))
    era = re.search(r'令和(\d+)年1月1日', str(rows[0][0]))
    if not era: return
    year = 2018 + int(era.group(1))
    if year <= published('jp'): return
    values = {'WD-JP-' + str(r[0])[:5]: int(r[5]) for r in rows if r[0] and str(r[0]).isdigit() and isinstance(r[5], (int, float))}
    if len(values) > 1500:
        apply('wd', values, year, 'jp', 'https://www.soumu.go.jp' + links[0], only=lambda n: n.startswith('JP-'))

def australia():
    """ABS estimated resident population at June 30 by SA2 (Regional population, table 32180DS0001), in wd/AU-*.json."""
    import openpyxl
    for year in range(NOW, published('au'), -1):
        edition = f'{year - 1}-{str(year)[2:]}'
        url = f'https://www.abs.gov.au/statistics/people/population/regional-population/{edition}/32180DS0001_{edition}.xlsx'
        try: data = get(url, binary=True)
        except urllib.error.URLError: continue
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
        values = {}
        for ws in wb.worksheets:
            if not ws.title.startswith('Table'): continue
            rows = ws.iter_rows(values_only=True)
            years = next((r for r in rows if r and year in r), None)
            if not years: continue
            col = list(years).index(year)
            for r in rows:
                if len(r) > col and isinstance(r[6], int) and isinstance(r[col], int): values[f'WD-AU-{r[6]}'] = r[col]
        if len(values) > 2000:
            apply('wd', values, year, 'au', url, only=lambda n: n.startswith('AU-'))
        return

for step in (eurostat, israel, brazil, united_states, canada, japan, australia):
    before = dict(datasets.get('_years') or {})
    try: step()
    except Exception as e: print(step.__name__, 'failed:', e)
    else:
        if (datasets.get('_years') or {}) == before: print(step.__name__ + ': nothing newer than what is published')
json.dump(manifest, open(manifest_path, 'w'), indent=2, ensure_ascii=False)
open(manifest_path, 'a').write('\n')
